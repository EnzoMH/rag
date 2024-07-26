from pptx.dml.color import RGBColor
from pptx.util import Pt

def set_font(text_frame, font_name, font_size, bold=False, color=None):
    """
    텍스트 프레임의 글꼴 이름, 크기, 굵기, 색상을 설정합니다.
    """
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = bold
            if color:
                run.font.color.rgb = color
